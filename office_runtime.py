"""Bounded Office artifact creation for the local Hermes workspace."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    Workbook = None
    Font = None
    PatternFill = None

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    Presentation = None
    Pt = None


OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".pdf", ".csv", ".md", ".txt"}
INVALID_FILENAME = re.compile(r"[<>:\"/\\|?*\x00-\x1f]")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


class OfficeRuntime:
    def __init__(self, output_root: Path | str, approved_roots: Iterable[Path | str]) -> None:
        self.output_root = Path(output_root).resolve()
        self.output_root.mkdir(parents=True, exist_ok=True)
        roots = [Path(root).resolve() for root in approved_roots if Path(root).exists()]
        roots.append(self.output_root)
        self.approved_roots = tuple(dict.fromkeys(roots))

    def capabilities(self) -> dict[str, bool]:
        return {
            "docx": Document is not None,
            "xlsx": Workbook is not None,
            "pptx": Presentation is not None,
        }

    @staticmethod
    def _title(value: Any) -> str:
        title = str(value or "").strip()
        if not title:
            raise ValueError("office_title_required")
        if len(title) > 200:
            raise ValueError("office_title_too_long")
        return title

    @staticmethod
    def _filename(title: str) -> str:
        cleaned = INVALID_FILENAME.sub("-", title).strip(" .")
        cleaned = re.sub(r"\s+", " ", cleaned)[:80].strip(" .")
        return cleaned or "Hermes artifact"

    def _target(self, title: str, extension: str) -> Path:
        base = self._filename(title)
        candidate = self.output_root / f"{base}{extension}"
        index = 2
        while candidate.exists():
            candidate = self.output_root / f"{base} ({index}){extension}"
            index += 1
        resolved = candidate.resolve()
        if resolved.parent != self.output_root:
            raise PermissionError("office_output_outside_workspace")
        return resolved

    @staticmethod
    def _result(path: Path, artifact_type: str, title: str) -> dict[str, Any]:
        stat = path.stat()
        return {
            "title": title,
            "format": artifact_type,
            "path": str(path),
            "size": stat.st_size,
            "createdAt": utc_now(),
        }

    @staticmethod
    def _atomic_save(target: Path, saver) -> None:
        temporary = target.with_suffix(target.suffix + ".tmp")
        try:
            saver(temporary)
            temporary.replace(target)
        finally:
            if temporary.exists():
                temporary.unlink()

    def inventory(self, limit: int = 100) -> dict[str, Any]:
        safe_limit = max(1, min(500, int(limit)))
        items: list[dict[str, Any]] = []
        seen: set[Path] = set()
        for root in self.approved_roots:
            try:
                candidates = root.rglob("*")
                for path in candidates:
                    if len(items) >= safe_limit:
                        break
                    try:
                        if not path.is_file() or path.suffix.casefold() not in OFFICE_EXTENSIONS:
                            continue
                        resolved = path.resolve()
                        if resolved in seen or not (resolved == root or root in resolved.parents):
                            continue
                        seen.add(resolved)
                        stat = resolved.stat()
                        items.append({
                            "name": resolved.name,
                            "path": str(resolved),
                            "format": resolved.suffix.casefold().lstrip("."),
                            "size": stat.st_size,
                            "modifiedAt": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                            "generated": self.output_root in resolved.parents,
                        })
                    except OSError:
                        continue
            except OSError:
                continue
            if len(items) >= safe_limit:
                break
        items.sort(key=lambda item: item["modifiedAt"], reverse=True)
        return {
            "items": items[:safe_limit],
            "count": min(len(items), safe_limit),
            "roots": [str(root) for root in self.approved_roots],
            "capabilities": self.capabilities(),
        }

    def create_document(self, args: dict[str, Any]) -> dict[str, Any]:
        if Document is None:
            raise RuntimeError("office_docx_library_unavailable")
        title = self._title(args.get("title"))
        blocks = args.get("blocks") if isinstance(args.get("blocks"), list) else []
        if not blocks or len(blocks) > 100:
            raise ValueError("office_document_blocks_invalid")
        document = Document()
        document.core_properties.title = title
        document.add_heading(title, level=0)
        for block in blocks:
            if not isinstance(block, dict):
                raise ValueError("office_document_block_invalid")
            kind = str(block.get("type") or "paragraph")
            text = str(block.get("text") or "").strip()
            if not text or len(text) > 8000:
                raise ValueError("office_document_block_invalid")
            if kind == "heading":
                document.add_heading(text, level=max(1, min(3, int(block.get("level") or 1))))
            elif kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            elif kind == "paragraph":
                document.add_paragraph(text)
            else:
                raise ValueError("office_document_block_type_invalid")
        target = self._target(title, ".docx")
        self._atomic_save(target, document.save)
        return self._result(target, "docx", title)

    def create_spreadsheet(self, args: dict[str, Any]) -> dict[str, Any]:
        if Workbook is None:
            raise RuntimeError("office_xlsx_library_unavailable")
        title = self._title(args.get("title"))
        columns = args.get("columns") if isinstance(args.get("columns"), list) else []
        rows = args.get("rows") if isinstance(args.get("rows"), list) else []
        if not columns or len(columns) > 50 or len(rows) > 1000:
            raise ValueError("office_spreadsheet_shape_invalid")
        clean_columns = [str(value).strip()[:200] for value in columns]
        if any(not value for value in clean_columns):
            raise ValueError("office_spreadsheet_column_invalid")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(clean_columns)
        for row in rows:
            if not isinstance(row, list) or len(row) > len(clean_columns):
                raise ValueError("office_spreadsheet_row_invalid")
            sheet.append([str(value)[:2000] if value is not None else "" for value in row] + [""] * (len(clean_columns) - len(row)))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="263B54")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(clean_columns, 1):
            values = [column, *(str(sheet.cell(row=row, column=index).value or "") for row in range(2, min(sheet.max_row, 100) + 1))]
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(48, max(10, max(map(len, values)) + 2))
        target = self._target(title, ".xlsx")
        self._atomic_save(target, workbook.save)
        return self._result(target, "xlsx", title)

    def create_presentation(self, args: dict[str, Any]) -> dict[str, Any]:
        if Presentation is None:
            raise RuntimeError("office_pptx_library_unavailable")
        title = self._title(args.get("title"))
        slides = args.get("slides") if isinstance(args.get("slides"), list) else []
        if not slides or len(slides) > 40:
            raise ValueError("office_presentation_slides_invalid")
        presentation = Presentation()
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = title
        if len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = "Created by Hermes Project 5"
        for slide_data in slides:
            if not isinstance(slide_data, dict):
                raise ValueError("office_presentation_slide_invalid")
            slide_title = str(slide_data.get("title") or "").strip()
            bullets = slide_data.get("bullets") if isinstance(slide_data.get("bullets"), list) else []
            if not slide_title or len(slide_title) > 300 or not bullets or len(bullets) > 12:
                raise ValueError("office_presentation_slide_invalid")
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = slide_title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, bullet in enumerate(bullets):
                text = str(bullet).strip()
                if not text or len(text) > 1200:
                    raise ValueError("office_presentation_bullet_invalid")
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.level = 0
                paragraph.font.size = Pt(22)
        target = self._target(title, ".pptx")
        self._atomic_save(target, presentation.save)
        return self._result(target, "pptx", title)
